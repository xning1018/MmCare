# -*- coding: utf-8 -*-
import os
import sys
import subprocess

try:
    import docx
except ImportError:
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-docx"])

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.oxml.ns import qn
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.table import WD_TABLE_ALIGNMENT


def set_run_font(run, font_name="微软雅黑", size_pt=11, bold=False, color_rgb=None):
    run.font.name = font_name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    run.font.size = Pt(size_pt)
    run.bold = bold
    if color_rgb:
        run.font.color.rgb = RGBColor(*color_rgb)


def add_heading_custom(doc, text, level=1):
    h = doc.add_heading("", level=level)
    run = h.add_run(text)
    sizes = {0: 20, 1: 15, 2: 13, 3: 11.5}
    colors = {
        0: (13, 148, 136),
        1: (15, 118, 110),
        2: (17, 94, 89),
        3: (30, 41, 59),
    }
    set_run_font(run, "微软雅黑", sizes.get(level, 12), bold=True, color_rgb=colors.get(level, (0, 0, 0)))
    return h


def add_para(doc, text="", bold_prefix="", bullet=False, color_rgb=None):
    p = doc.add_paragraph(style="List Bullet" if bullet else None)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.35
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        set_run_font(r_pre, "微软雅黑", 10.5, bold=True, color_rgb=color_rgb or (15, 23, 42))
    if text:
        r_txt = p.add_run(text)
        set_run_font(r_txt, "微软雅黑", 10.5, bold=False, color_rgb=color_rgb or (51, 65, 85))
    return p


def add_styled_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    hdr_cells = table.rows[0].cells
    for i, h_text in enumerate(headers):
        hdr_cells[i].text = ""
        p = hdr_cells[i].paragraphs[0]
        r = p.add_run(h_text)
        set_run_font(r, "微软雅黑", 10, bold=True, color_rgb=(15, 118, 110))

    for row_data in rows:
        row_cells = table.add_row().cells
        for i, val in enumerate(row_data):
            row_cells[i].text = ""
            p = row_cells[i].paragraphs[0]
            r = p.add_run(val)
            set_run_font(r, "微软雅黑", 9.5, bold=(i == 0), color_rgb=(30, 41, 59))
    doc.add_paragraph()


def build_manual():
    doc = Document()

    # 全局默认字体
    doc.styles["Normal"].font.name = "微软雅黑"
    doc.styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    doc.styles["Normal"].font.size = Pt(10.5)

    # 标题
    t = add_heading_custom(doc, "MMCare 毫米波双雷达端侧 AI 智能康养监护系统", level=0)
    t.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    sub = doc.add_paragraph()
    sub.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    r_sub = sub.add_run("完整工程说明 · 日常使用手册 · 真机数据采集与 AI 训练指南 · 关键注意事项")
    set_run_font(r_sub, "微软雅黑", 11.5, bold=True, color_rgb=(100, 116, 139))

    # ==================== 第一部分 ====================
    add_heading_custom(doc, "一、 系统总体架构与硬件外设接线表", level=1)
    add_para(
        doc,
        "本系统基于瑞萨 RA6M5 (R7FA6M5BH3CFC, 200MHz Cortex-M33 内核 + 硬件单精度 FPU + 512KB SRAM) 设计，融合 HLK-LD2450 (24GHz 多目标轨迹雷达) 与 HLK-LD6002C (60GHz 三维点云跌倒雷达)，在单片机端侧运行 1024 粒子蒙特卡洛空间滤波、64 点微多普勒 STFT 频谱分析与 PINN-CfC 物理内嵌连续时间液态神经网络，并通过 ML307R 4G Cat.1 模组实现毫秒级云端大屏、微信公众号与 QQ 邮箱并发告警。",
    )

    add_heading_custom(doc, "1.1 硬件外设与引脚分配一览表", level=2)
    add_styled_table(
        doc,
        ["硬件模块", "功能角色", "单片机外设 / 引脚", "波特率 / 参数", "对应驱动源码"],
        [
            ["HLK-LD2450", "24GHz 2D 多目标轨迹雷达（AI 通道 A）", "UART7 (P401 TX / P402 RX)", "256000 bps", "API/radar_ld2450.c"],
            ["HLK-LD6002C", "60GHz 3D 点云跌倒雷达（AI 通道 B）", "UART9 (P602 TX / P601 RX)", "115200 bps", "API/radar_uart.c"],
            ["ML307R 4G Cat.1", "4G 蜂窝网络物联网上云模组", "UART4 (P502 TX / P501 RX)", "115200 bps", "API/ml307r.c"],
            ["OLED 显示屏", "本地实时显示距离、速度、温度与报警状态", "I2C2 (P512 SCL / P511 SDA)", "硬件 I2C (0x3C)", "API/oled.c / radar_display.c"],
            ["NTC 热敏电阻", "室内环境温度采集", "ADC0 AN000 (P000)", "12-bit ADC", "API/ntc_temp.c"],
            ["环境光传感器", "室内光照检测与天黑有人自动夜灯联动", "GPIO (P014 DO / P100 LED)", "高低电平触发", "API/light_sensor.c"],
            ["MY1680 语音模块", "跌倒现场真人语音播报告警", "UART 串口控制", "语音播报", "API/my1680.c"],
            ["无源蜂鸣器+按键", "本地 2kHz 高分贝声学报警 + 手动消警按键", "GPT3 PWM (P403) + 按键 (P004)", "2kHz PWM", "API/fall_alarm.c / buzzer.c"],
        ],
    )

    add_heading_custom(doc, "1.2 核心代码文件结构说明", level=2)
    add_para(doc, "src/hal_entry.c：系统主入口、前后台非阻塞轮询调度、5 秒冷启动热机防误报控制、4G 上报参数组装。", bold_prefix="• ", bullet=True)
    add_para(doc, "API/ai_fall_engine.c / .h：端侧双雷达 AI 跌倒检测核心引擎（1024 粒子滤波、16 维物理流形提取、64 点 STFT、PINN-CfC 液态网络前向推理、双通道 111/1/11 协同判决）。", bold_prefix="• ", bullet=True)
    add_para(doc, "API/ai_weights.h：Flash 对称 INT8 量化存储的预训练神经网络权重库与类别原型矩阵（上电自动反量化至 SRAM FP32 张量池）。", bold_prefix="• ", bullet=True)
    add_para(doc, "API/radar_ld2450.c / .h：HLK-LD2450 24GHz 雷达 256000bps 串口协议帧解析驱动（提取多目标二维坐标 X/Y 与径向速度 V）。", bold_prefix="• ", bullet=True)
    add_para(doc, "API/radar_uart.c / .h：HLK-LD6002C 60GHz 雷达 115200bps TF 协议帧解析驱动（提取三维点云簇 X/Y/Z、垂直下坠速度、三维冲击加速度、点云宽高比及 0x0E02 跌倒帧）。", bold_prefix="• ", bullet=True)
    add_para(doc, "API/ml307r.c / .h：ML307R 4G Cat.1 模组 AT 状态机驱动，支持日常 5 秒心跳与跌倒 0ms 抢占式极速上报。", bold_prefix="• ", bullet=True)
    add_para(doc, "dashboard.html：智能康养监护终端实时遥测 Web 大屏（0.3s 极速轮询、雷达扫描动画、触发代码 111/1/11 显示、7 秒冷却微信与 QQ 邮箱并发推送）。", bold_prefix="• ", bullet=True)
    add_para(doc, "scripts/collect_and_train.py 与 启动AI姿态采集与训练.bat：电脑端真机雷达数据采集、PINN-CfC 神经网络重训练与 C 语言 INT8 权重头文件自动生成器。", bold_prefix="• ", bullet=True)

    # ==================== 第二部分 ====================
    add_heading_custom(doc, "二、 核心工程逻辑与双雷达多参数融合跌倒判决机制", level=1)

    add_heading_custom(doc, "2.1 主循环调度与 5 秒开机防误报机制", level=2)
    add_para(doc, "上电/复位 5 秒冷启动保护：", "1. ")
    add_para(doc, "开机或按复位键后的前 5 秒（5000ms）为硬件冷启动热机期。在此期间：1024 粒子滤波器在收到第一帧有效人体坐标时直接将全部粒子锚定在真实位置周围（消除粒子从默认中心滑向人体产生的虚假高速收敛冲击），同时等待 16 点滑动窗口完全填满，前 5 秒内强制屏蔽一切报警输出，彻底解决开机/复位误报问题。")
    add_para(doc, "每 100ms（10Hz）实时 AI 推理：", "2. ")
    add_para(doc, "后台通过 AGT0 硬件定时器维护 1ms 系统时基 g_system_tick，每 100ms 调用一次 AI_FallEngine_Step()，同时运行通道 A（LD2450）和通道 B（LD6002C）的 16 维物理特征提取与 PINN-CfC 液态神经网络推理。")
    add_para(doc, "4G Cat.1 双速上报机制：", "3. ")
    add_para(doc, "平时正常状态下每 5 秒上报一次温湿度、距离、速度与姿态；一旦检测到跌倒上升沿（combined_fall_code 非 0），0ms 立即打断常规等待，抢占发送 /api/v1/fall/report 紧急跌倒事件与 /api/data 状态包。")

    add_heading_custom(doc, "2.2 双雷达各自内部的多参数物理融合（通道内融合）", level=2)
    add_para(doc, "本系统绝非只看单一速度或单一雷达，而是将 LD2450 与 LD6002C 各自的多维物理量提取为同构的 16 维物理流形特征送入 PINN-CfC 液态神经网络（4 分类：Class 0 正常慢走、Class 1 快走/急转身、Class 2 弯腰/下蹲/坐下、Class 3 失衡跌倒报警）：")

    add_styled_table(
        doc,
        ["特征维度", "代码变量名", "通道 A (LD2450 24GHz) 物理含义", "通道 B (LD6002C 60GHz 3D点云) 物理融合含义"],
        [
            ["feat[0]", "mean_r", "16点窗口内人体距雷达平均二维径向距离 (/6.0m)", "16点窗口内人体平均三维空间径向距离 sqrt(X²+Y²+Z²) (/6.0m)"],
            ["feat[1]", "span_r", "窗口内水平最大位移跨度 (/2.5m)", "三维径向跨度 + 相对自标定站立基线的垂直高度坍塌率 rel_y_shift"],
            ["feat[2]", "mean_v", "多普勒速度(60%) + 坐标导数速度(40%) 的平均速度", "三维移动速度(50%) + 垂直下坠速度 -dH/dt(30%) + 径向速度(20%) 均值"],
            ["feat[3]", "max_v", "窗口内瞬时峰值径向速度 (/2.5m/s)", "窗口内瞬时峰值三维合成下坠速度 (/2.5m/s)"],
            ["feat[4]", "mean_v_early", "前半窗口（前 0.8 秒）平均速度", "前半窗口（跌倒失衡阶段）平均三维合成速度"],
            ["feat[5]", "mean_v_late", "后半窗口（后 0.8 秒）平均速度", "后半窗口（倒地阶段）平均三维合成速度"],
            ["feat[6]", "max_acc", "瞬时峰值加速度 max|dv/dt| (/6.0m/s²)", "三维点云瞬时峰值冲击加速度 (/6.0m/s²)"],
            ["feat[7]", "max_jerk", "瞬时峰值加加速度（冲击度 Jerk, /25.0m/s³）", "三维点云瞬时峰值加加速度 Jerk (/25.0m/s³)"],
            ["feat[8]", "zero_crossings", "速度方向过零率（区分原地晃动与单向倒地）", "垂直下坠速度方向过零率"],
            ["feat[9]", "collapse_ratio", "速度骤降坍塌比 (v前 - v后)/(v前 + 0.15)", "三维下坠速度骤降坍塌比（前半段急坠 -> 后半段骤停）"],
            ["feat[10]", "post_stillness", "跌后二维位置静止度 exp(-σ_r * 8.0)", "跌后三维点云空间静止度 exp(-σ_r3d * 8.0)"],
            ["feat[11]", "pinn_impact", "物理冲击指数 max_v * max_jerk * post_stillness", "物理冲击指数 × 三维点云卧姿宽高比形变增益 rel_aspect_gain"],
            ["feat[12~15]", "bands[0~3]", "64点汉宁窗微多普勒 FFT 四个频段能量占比", "64点三维下坠速度序列 FFT 四个频段能量占比"],
        ],
    )

    add_heading_custom(doc, "2.3 双雷达之间的协同判决与触发代码（111 / 1 / 11）", level=2)
    add_para(doc, "由于 LD2450 覆盖范围广（0.5m~6m，水平 ±60°），而 LD6002C 的 60GHz 三维点云在近距离核心区（0.5m~2.5m）对高度落差和卧姿宽高比极其敏锐，系统采用“双通道独立并行或门触发 + 双雷达共振分级编码”：")

    add_styled_table(
        doc,
        ["触发代码", "触发通道", "具体数值判定标准（需开机满 5 秒后）"],
        [
            [
                "111",
                "通道 A (LD2450 24GHz)\n独立触发报警",
                "16 点轨迹滑窗填满后，满足任一分支即报警：\n"
                "① AI + 运动学守恒：AI 跌倒概率 probs[3] ≥ 56%，且峰值速度 feat[3] ≥ 0.42 (≥1.05m/s)，且跌后静止度 feat[10] ≥ 0.40，且物理残差 < 2.2；\n"
                "② PINN 强冲击直通：峰值速度 feat[3] ≥ 0.55 (≥1.38m/s)，且峰值加速度 feat[6] ≥ 0.35，且速度坍塌比 feat[9] ≥ 0.32，且跌后静止度 feat[10] ≥ 0.52。",
            ],
            [
                "1",
                "通道 B (LD6002C 60GHz)\n独立触发报警",
                "满足任一分支即报警：\n"
                "① 底层 0x0E02 跌倒帧：雷达原生上报 is_fall == 1 且 is_human != 0；\n"
                "② 3D 点云多参数 + AI 推理：点云窗口满 10 点后，AI 跌倒概率 probs_60[3] ≥ 56%，且峰值三维合成速度 max_v ≥ 0.38 (≥0.95m/s)，且跌后三维静止度 post_stillness ≥ 0.42，且速度坍塌比 collapse_ratio ≥ 0.25。",
            ],
            [
                "11",
                "LD2450 + LD6002C\n双雷达同时共振触发",
                "通道 A 与通道 B 在同一时刻均处于跌倒锁存状态（轨迹多普勒冲击 + 三维点云高度坍塌/卧姿宽高比同时确认，最高置信度）。",
            ],
        ],
    )

    add_para(doc, "报警保持与自动恢复机制：", bold_prefix="• ", bullet=True)
    add_para(doc, "一旦触发跌倒，声光报警与云端报警状态最短保持 8 秒（AI_FALL_HOLD_MS = 8000ms）。若满 3 秒后，雷达检测到人体已重新站起并恢复正常行走（移动速度 > 0.35~0.48 m/s 且 AI 分类回到慢走/快走 Class 0/1），系统将提前自动解除报警。")

    # ==================== 第三部分 ====================
    add_heading_custom(doc, "三、 日常使用操作手册（硬件 + 网页大屏 + 微信/QQ 邮箱）", level=1)

    add_heading_custom(doc, "步骤 1：设备安装、上电与自适应高度标定", level=2)
    add_para(doc, "将开发板与双雷达接通 5V 电源（或烧录固件后按复位键）。", bold_prefix="1. ")
    add_para(doc, "等待 5 秒冷启动期：上电前 5 秒系统自动锁定不报警，完成粒子滤波首帧锚定。", bold_prefix="2. ")
    add_para(doc, "正常走动 3~5 秒完成自标定：上电后人在雷达前方 1m~3m 处正常走两步。LD6002C 会自动学习并锁定当前安装高度下的常态站立垂直基线（s_pc60_base_y_m）与站立点云宽高比基线（s_pc60_base_aspect）；LD2450 也会自动标定 32×32 房间网格过滤静态家具回波。", bold_prefix="3. ")

    add_heading_custom(doc, "步骤 2：打开实时监控大屏（本地或云端任选其一）", level=2)
    add_para(doc, "方式 A（推荐，直接打开电脑本地网页）：双击工程目录下的 dashboard.html，浏览器会自动以 0.3 秒（300ms）极速轮询连接云服务器 http://182.254.135.92:8080，实时显示温度、距离、速度、姿态及 [触发代码: 111 / 1 / 11]。", bold_prefix="• ", bullet=True)
    add_para(doc, "方式 B（访问云服务器网页 http://182.254.135.92:8000）：若更新过本地 dashboard.html，请将工程目录下 deploy_cmd.txt 中的整行命令复制到腾讯云服务器（182.254.135.92）终端执行，随后在浏览器按 Ctrl + F5 强制刷新。", bold_prefix="• ", bullet=True)
    add_para(doc, "重要提醒：请确保浏览器中只打开一个监控大屏标签页（关闭多余的旧标签页），因为微信与 QQ 邮箱推送是由当前打开的网页标签页实时发出的。", bold_prefix="⚠️ ", bullet=True, color_rgb=(185, 28, 28))

    add_heading_custom(doc, "步骤 3：触发跌倒报警与查看微信 / QQ 邮箱推送", level=2)
    add_para(doc, "当系统检测到跌倒时：", bold_prefix="1. ")
    add_para(doc, "开发板本地：OLED 屏幕显示 FALL!，蜂鸣器高频鸣叫，语音模块播报跌倒告警（按下开发板 P004 按键可手动静音消警）。", bullet=True)
    add_para(doc, "监控大屏：顶部弹出单行醒目红色横幅“【紧急跌倒警报】检测到长者跌倒！[触发代码: 111]”（或 1 / 11），状态卡片变红并鸣响提示音。", bullet=True)
    add_para(doc, "微信与 QQ 邮箱：0ms 并发推送含 [触发代码: 111 / 1 / 11]、告警时间、距离、速度与环境温度的通知卡片。防刷屏冷却时间为 7 秒（两次推送间隔满 7 秒即可再次推送）。", bullet=True)
    add_para(doc, "手动测试推送通道：点击网页右上角“📲 微信推送测试”按钮，可随时验证微信公众号与 QQ 邮箱弹窗是否正常。", bold_prefix="2. ")

    # ==================== 第四部分 ====================
    add_heading_custom(doc, "四、 真机数据采集与 AI 模型重训练完整操作指南（图形化倒计时界面）", level=1)
    add_para(
        doc,
        "本工程为您专门配备了全中文图形化倒计时引导式采集与训练中心（双击启动_AI采集与训练界面.bat / 启动AI姿态采集与训练.bat，底层调用 scripts/collect_and_train.py）。无需额外接串口线，只要开发板上电连上 4G，电脑即可通过云端实时拉取雷达数据进行采集！",
    )

    add_heading_custom(doc, "4.1 第一步：启动图形化采集中心并确认开发板在线", level=2)
    add_para(doc, "确保 RA6M5 开发板已上电（上电满 5 秒后进入正常工作状态）。", bold_prefix="1. ")
    add_para(doc, "在电脑工程文件夹下，双击运行【启动AI姿态采集与训练.bat】（或【双击启动_AI采集与训练界面.bat】）。", bold_prefix="2. ")
    add_para(doc, "此时会弹出深青色图形界面“MMCare 瑞萨 RA6M5 · 分步倒计时引导式 AI 姿态采集与训练中心”。查看窗口右上角状态栏，显示“📡 板端4G: 在线 | 实时距离: X.XXm | 速度: X.Xcm/s”即代表连接成功！", bold_prefix="3. ")

    add_heading_custom(doc, "4.2 第二步：按顺序点击 4 个按钮录制真机动作（带 3 秒准备语音提示）", level=2)
    add_para(doc, "界面底部共有 ①、②、③、④ 四个彩色采集按钮。每次点击任意按钮后，系统都会先给您 3 秒准备倒计时（伴随“滴-滴-滴”提示音），让您有充足时间从电脑旁走到雷达前；当听到一声长音“滴——”时，正式开始做动作：", bold_prefix="")
    add_para(doc, "点击【① 采集 正常慢走 (倒计时 25 秒)】：听到长音后，在雷达前 1m~4m 区域内正常慢走、来回散步 25 秒（采集日常平稳慢走特征 Class 0）。", bold_prefix="• ", bullet=True)
    add_para(doc, "点击【② 采集 快走 & 急转身 (倒计时 25 秒)】：听到长音后，在雷达前做快步走过、走到一半突然急转身、挥手等动作 25 秒（专门训练 AI 认识快走与急停，防止把快走误报为跌倒 Class 1）。", bold_prefix="• ", bullet=True)
    add_para(doc, "点击【③ 采集 弯腰 & 下蹲 (倒计时 25 秒)】：听到长音后，在雷达前反复做原地弯腰捡东西、慢慢蹲下停顿 2 秒再站起来的动作 25 秒（专门训练 AI 认识弯腰下蹲，防止把蹲下误报为跌倒 Class 2）。", bold_prefix="• ", bullet=True)
    add_para(doc, "点击【④ 录制 1 次 真实摔倒 (倒计时 6 秒/次)】：听到长音后，立即在软垫上做出快速摔倒/快速深蹲到底的动作，并保持躺地/蹲地静止不动 3~4 秒，直到听到两声“滴-滴”完成音！（建议重复点击按钮 ④ 录制 5 ~ 10 次不同位置或方向的摔倒样本 Class 3）。", bold_prefix="• ", bullet=True)
    add_para(doc, "防呆与纠错功能：如果在某次倒计时期间没准备好，随时点击界面中间的【⏹️ 停止当前倒计时】即可取消本次录制，不会保存脏数据；若想全部重录，点击【🗑️ 清空真机样本】即可。", bold_prefix="💡 ", bullet=True)

    add_heading_custom(doc, "4.3 第三步：一键训练并烧录固化到单片机", level=2)
    add_para(doc, "四个动作采集完成后，直接点击界面最底部的绿色大按钮：【🚀 一键训练 PINN-CfC 液态神经网络 并自动更新 C 权重头文件 (API/ai_weights.h)】。", bold_prefix="1. ")
    add_para(doc, "等待约 3~5 秒，界面中央会显示“🎉 INT8量化后独立验证集准确率 XX.XX% | 跌倒召回率 XX% | 误报率 XX%”，此时您采集的真机数据已经自动转换成 C 语言 INT8 量化权重数组并写入了 API/ai_weights.h！", bold_prefix="2. ")
    add_para(doc, "打开 Keil MDK 工程（Mmcare.uvprojx），按 F7 (Build) 重新编译，再按 F8 (Download) 烧录进 RA6M5 开发板，您的专属实测 AI 模型即刻在单片机上永久生效！", bold_prefix="3. ")

    # ==================== 第五部分 ====================
    add_heading_custom(doc, "五、 关键注意事项与常见问题排查（避坑必看）", level=1)

    add_para(doc, "1. 模拟跌倒测试的正确动作要领（“快倒 + 原地静止 1.5~2 秒”）：", bold_prefix="")
    add_para(
        doc,
        "无论是 LD2450（通道 A）还是 LD6002C（通道 B），AI 引擎为了滤除日常走动和弯腰误报，都会同时检验【瞬时下坠/冲击速度 max_v】、【前后半窗速度骤降坍塌比 collapse_ratio】以及【跌后静止度 post_stillness】。如果您只是快速晃一下身体就立刻站起来走开，后半窗不静止，AI 会将其识别为“快步走/急转身 (Class 1)”而拦截报警。因此在测试跌倒时，请先做出快速倒下/快速深蹲动作（激发速度、高度差与卧姿宽高比），随后在原地保持静止不动 1.5 ~ 2 秒，即可稳定触发报警！",
        bullet=True,
    )

    add_para(doc, "2. 微信与 QQ 邮箱 7 秒防刷屏冷却：", bold_prefix="")
    add_para(
        doc,
        "为防止 PushPlus 接口因高频连续请求触发限流，网页端设置了 7 秒（7000ms）推送冷却。连续测试跌倒时，请等待上次报警解除并间隔满 7 秒后再做下一次跌倒动作。",
        bullet=True,
    )

    add_para(doc, "3. 切勿在浏览器同时打开多个监控大屏页面：", bold_prefix="")
    add_para(
        doc,
        "微信与邮件推送由浏览器打开的网页前端实时触发。若同时开着旧版和新版两个标签页，旧页面会在后台抢发旧格式推送。请务必只保留一个最新大屏页面。",
        bullet=True,
    )

    add_para(doc, "4. Keil MDK 32KB 编译体积红线：", bold_prefix="")
    add_para(
        doc,
        "若您的 Keil MDK 处于 32KB 评估版授权状态，总代码体积（Code + RO-data + RW-data）上限为 32,768 Bytes。当前固件经高度精简优化后体积为 32,472 Bytes（余量 296 字节），0 Error(s), 0 Warning(s)。后续若自行修改 C 代码，切勿引入 stdio.h 中的 printf / snprintf 等大体积标准库函数。",
        bullet=True,
    )

    add_para(doc, "5. 稳定版备份与一键回滚说明：", bold_prefix="")
    add_para(
        doc,
        "当前稳定版本已完整备份在工程目录下的 Backup_20261012_AI_111_7s/ 文件夹、dashboard_backup_ai_111_7s_20261012.html 以及 Git 提交节点 567d532，任何时候均可一键恢复。",
        bullet=True,
    )

    desktop_path = r"C:\Users\ls\Desktop\MMCare_完整工程说明与使用手册.docx"
    project_path = r"c:\Users\ls\RASmartConfigurator\Mmcare\MMCare_完整工程说明与使用手册.docx"
    doc.save(desktop_path)
    doc.save(project_path)
    print("Saved Word document to:", desktop_path)
    print("Saved backup copy to:", project_path)


if __name__ == "__main__":
    build_manual()
