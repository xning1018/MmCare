from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

prs = Presentation()

# Title slide
slide_layout = prs.slide_layouts[0]
slide = prs.slides.add_slide(slide_layout)
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Mmcare 智能看护与跌倒检测系统"
subtitle.text = "基于瑞萨 RA6M5 的非阻塞式物联网架构\n\n毕业设计 / 项目答辩"

def add_bullet_slide(prs, title_text, bullets):
    slide_layout = prs.slide_layouts[1]
    slide = prs.slides.add_slide(slide_layout)
    title = slide.shapes.title
    title.text = title_text
    tf = slide.placeholders[1].text_frame
    tf.text = bullets[0]
    for bullet in bullets[1:]:
        p = tf.add_paragraph()
        p.text = bullet
        p.level = 0
    return slide

add_bullet_slide(prs, "1. 项目背景与需求分析", [
    "人口老龄化加剧，老年人独居跌倒成为重大安全隐患。",
    "传统摄像头方案存在严重的隐私泄露风险，且受光照影响大。",
    "需求分析：",
    " - 需要高精度、全天候的非接触式跌倒检测（毫米波雷达）。",
    " - 需要实时将报警信息推送到云端（4G 蜂窝网络）。",
    " - 异常状态需要本地声光报警及语音提示（蜂鸣器、LED、语音模块）。",
    " - 系统需具备高可靠性，不能因网络波动卡死业务逻辑。"
])

add_bullet_slide(prs, "2. 系统整体架构设计", [
    "主控芯片：瑞萨电子 RA6M5 高性能 MCU（基于 FSP 库）。",
    "感知层：",
    " - LD2450 轨迹雷达：定位人员坐标 (X,Y) 及移动速度 (V)。",
    " - LD6002C 跌倒雷达：专业姿态与跌倒识别，输出跌倒状态与距离。",
    " - 环境感知：NTC 测温 + 光敏电阻检测环境光强。",
    "交互层：",
    " - OLED 屏幕（状态显示）、RGB 指示灯、蜂鸣器、MY1680 语音芯片。",
    "通信层：",
    " - ML307R Cat.1 4G 通信模组，通过 HTTP POST 协议对接腾讯云/自建轻量服务器。"
])

add_bullet_slide(prs, "3. 核心功能一：雷达数据融合与分析", [
    "多维度雷达协同探测：",
    " - 轨迹跟踪：解析 LD2450 串口协议，提取目标 X/Y/V 绝对坐标及速度。",
    " - 异常跳变检测：实时比对帧间位移 (Delta) 与速度阈值，对突变进行过滤和标记。",
    " - 跌倒判决：融合 LD6002C 跌倒算法输出，结合轨迹异常标志，实现高置信度的跌倒报警 (状态字: 11)。",
    " - 目标丢失保护：3秒无目标自动清空历史缓存，防止幽灵目标。"
])

add_bullet_slide(prs, "4. 核心功能二：智能场景与声光联动", [
    "跌倒紧急响应：",
    " - 跌倒状态触发时，蜂鸣器启动 500ms 周期的高频翻转报警 (滴-滴-滴)。",
    " - 触发 MY1680 语音芯片播放专属报警/求救语音。",
    " - 系统指示灯立即切换至紧急颜色 (特定频闪或红灯)。",
    "智能夜灯控制 (多传感器融合)：",
    " - 条件判断：光敏传感器检测到环境无光 (Status == 1)。",
    " - 联动雷达：若同时雷达检测到人类存在且最近3秒内有移动，自动点亮白光 LED 灯带。",
    "断电保护展示：优先展示最高等级的硬件状态指示。"
])

add_bullet_slide(prs, "5. 核心功能三：云端互联与 4G 传输", [
    "ML307R 模组 AT 指令封装：",
    " - 自动完成 CPIN 检查、驻网 (CEREG) 检查及 PDP 拨号。",
    " - 纯底层 TCP 透传重构 HTTP POST，彻底解决传统模块底层死锁问题。",
    "JSON 数据序列化上报：",
    " - 自动组装 device_id, temp, human, fall, dist, x, y, v。",
    " - 为规避单片机浮点处理开销，自动执行浮点转整型映射 (如 25.5°C 转换为 255)。",
    "双向控制链路：",
    " - 解析服务器下发的 HTTP 200 OK 报文，提取 JSON 指令 (如 LED_ON) 反控本地设备。"
])

add_bullet_slide(prs, "6. 软件架构亮点一：非阻塞与事件驱动", [
    "摒弃死等，拥抱并发 (System_Yield)：",
    " - 难点：传统的 4G HTTP 发送需要耗时数秒（连接基站、发报文、等回复），导致主循环卡死，进而漏报雷达数据。",
    " - 解决：在 4G 的每一个超时等待循环中嵌入 System_Yield() 后台维护函数。",
    " - 成果：即使网络正在缓慢拨号，雷达数据的解析、蜂鸣器的定时鸣叫、LED 的状态切换依然在后台完美并发运行，互不干扰！"
])

add_bullet_slide(prs, "7. 软件架构亮点二：低功耗与按需上报机制", [
    "休眠与唤醒的极致压榨：",
    " - MCU 休眠：基于瑞萨 FSP 的 R_LPM_LowPowerModeEnter，在无处理任务时关闭 CPU 时钟节省功耗。",
    " - 精准时基：利用 AGT 定时器提供每 10ms 的精准 g_system_tick 唤醒系统。",
    "数据阈值过滤算法 (按需联网)：",
    " - 摒弃了无脑的 8 秒定时上传机制。",
    " - 采用 Delta 空间算法比对：只有当人体位移跨度大于 50 毫米，或发生严重跌倒时，才触发 4G 模块发射数据。",
    " - 极大降低了设备的无效心跳流量与功耗发射，为纯电池供电模式奠定理论基础。"
])

add_bullet_slide(prs, "8. 总结与展望", [
    "项目总结：",
    " - 成功实现了一套完整的、无感知的智能看护系统硬件与固件生态。",
    " - 巧妙地克服了单线程裸机环境下的网络阻塞难题，实现了高可靠声光并发。",
    " - 代码架构清晰规范，核心部分采用了真正的事件驱动与按需上报机制。",
    "未来展望：",
    " - 架构升级：引入 MQTT 协议支持，实现更低延迟的云端持久双向长连接。",
    " - 算法优化：优化雷达空间算法，加入多目标并发跟踪和跌倒姿势精细化识别。",
    " - 兜底机制：增加 1 小时级别的心跳保活 (Heartbeat) 逻辑，完善系统断线重连监控体系。"
])

prs.save(r'E:\Mmcare智能看护系统答辩.pptx')
print("PPT successfully generated in E:\\Mmcare智能看护系统答辩.pptx")
